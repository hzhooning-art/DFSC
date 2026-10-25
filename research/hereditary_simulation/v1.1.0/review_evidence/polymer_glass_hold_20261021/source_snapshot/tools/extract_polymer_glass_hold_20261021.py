"""Read source observations only; do not touch source fitted columns or modify XLSX."""
from pathlib import Path
import hashlib,json
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
SPEC=ROOT/'experiments/specs/polymer_glass_hold_20261021.json'
OUT=ROOT/'results/polymer_glass_hold_20261021'
def main():
    spec=json.loads(SPEC.read_text());p=ROOT/'data/external/polymer_glass_20261021'/spec['workbook']
    assert hashlib.sha256(p.read_bytes()).hexdigest()==spec['workbook_sha256']
    OUT.mkdir(parents=True,exist_ok=True)
    if (OUT/'extracted.json').exists():raise FileExistsError('Preserve extracted source')
    # Protocol is saved before response extraction, not before public data publication.
    (OUT/'extraction_protocol_receipt.json').write_text(json.dumps({
        'spec_sha256':hashlib.sha256(SPEC.read_bytes()).hexdigest(),
        'extractor_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        'source_sha256':spec['workbook_sha256'],'frozen_before_response_extraction':True,
        'externally_blinded':False},indent=2),encoding='utf-8')
    w=openpyxl.load_workbook(p,read_only=True,data_only=True);s=w[spec['sheet']];traces=[]
    for temp,(xc,yc) in zip(spec['temperature_K'],spec['experimental_column_pairs_1_based']):
        assert s.cell(3,xc).value==f'{temp}K'
        assert s.cell(4,xc).value=='Time (s)' and s.cell(4,yc).value=='E(t) (MPa)'
        rows=[]
        for i in range(spec['source_rows_start'],s.max_row+1):
            t,y=s.cell(i,xc).value,s.cell(i,yc).value
            if t is None and y is None:continue
            if not isinstance(t,(int,float)) or not isinstance(y,(int,float)):raise ValueError('Unpaired/nonnumeric measured source cells')
            rows.append({'source_row':i,'time_seconds':float(t),'modulus_MPa':float(y)})
        assert len(rows)>=10 and all(b['time_seconds']>a['time_seconds'] for a,b in zip(rows,rows[1:]))
        traces.append({'temperature_K':temp,'source_columns':[xc,yc],'records':rows})
    w.close()
    (OUT/'extracted.json').write_text(json.dumps({'sheet':spec['sheet'],'traces':traces,
        'source_workbook_unchanged':True,'source_fitted_columns_used':False},indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'curves':len(traces),'experimental_rows':[len(t['records']) for t in traces],
        'time_ranges':[[t['records'][0]['time_seconds'],t['records'][-1]['time_seconds']] for t in traces],
        'response_values_not_printed':True},indent=2))
if __name__=='__main__':main()
